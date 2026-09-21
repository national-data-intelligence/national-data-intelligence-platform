#!/usr/bin/env python3
"""Build the sample deck and walkthrough that ship on the NHDW shelf.

Not downloaded and not lorem: both describe the actual system, so what a
presenter opens in the console looks like something somebody would really show.
The reference PDFs on that shelf are the project's own documents and live in the
NHDW repository — the seeder takes a --docs directory for those.

    python3 tools/make-sample-media.py

Needs python-pptx and Pillow, and ffmpeg for the encode. Run it on a machine
that has them; the results are committed, so nobody else has to.
"""
from __future__ import annotations

import pathlib
import shutil
import subprocess
import sys
import tempfile

HERE = pathlib.Path(__file__).resolve().parent.parent
OUT = HERE / "samples"

SLIDES = [
    ("National Health Data Warehouse", "Ministry of Health · Ethiopia", [
        "One governed warehouse behind five dashboards",
        "Data Vault 2.0 — history you can reconstruct, not overwrite",
        "The same access rules on every read path"]),
    ("Where the data comes from", "Seven ingestors, landing the source as it is", [
        "DHIS2 — aggregate programme reporting",
        "eLMIS — 7.5M requisition lines of commodity history",
        "mBrana, EthioEMR, MFR, climate, and the facility registry",
        "Raw stays all-string; typing happens downstream, on purpose"]),
    ("The vault", "Hubs, links and satellites on conformed keys", [
        "A facility is one facility across six source systems",
        "Same-as links keep the source system in the key, so nothing collapses",
        "Steward decisions are durable data the build reads, not a live edit"]),
    ("The marts", "The numbers every dashboard reads", [
        "bi_quality — completeness, timeliness, the incident queue",
        "bi_supply — availability, stockouts, months of stock, order fill",
        "bi_service, bi_public, bi_mdm"]),
    ("Governed access", "One decision point, every path", [
        "OPA decides; the gateway enforces — including the SQL editor and chat",
        "Row filters by scope, column masking by classification",
        "Fail closed: an unclassified sensitive column is caught, not revealed"]),
    ("What you are about to see", "A short tour", [
        "Internal BI — the operational picture",
        "Public portal — what a citizen sees",
        "Admin — quality, topology, people and policy"]),
]

INK, DIM, ACC = (0x14, 0x16, 0x1B), (0x58, 0x5E, 0x6C), (0x0F, 0x7F, 0x57)


def build_deck() -> None:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt
    except ImportError:
        sys.exit("the deck needs python-pptx:  pip install python-pptx")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for title, sub, bullets in SLIDES:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        tb = s.shapes.add_textbox(Inches(.9), Inches(.8), Inches(11.5), Inches(1.4)).text_frame
        tb.text = title
        r = tb.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.color.rgb = Pt(40), True, RGBColor(*INK)
        sb = s.shapes.add_textbox(Inches(.9), Inches(1.9), Inches(11.5), Inches(.7)).text_frame
        sb.text = sub
        r = sb.paragraphs[0].runs[0]
        r.font.size, r.font.color.rgb = Pt(19), RGBColor(*ACC)
        bb = s.shapes.add_textbox(Inches(.9), Inches(2.9), Inches(11.5), Inches(4)).text_frame
        bb.word_wrap = True
        for i, b in enumerate(bullets):
            p = bb.paragraphs[0] if i == 0 else bb.add_paragraph()
            p.text = "— " + b
            p.runs[0].font.size = Pt(18)
            p.runs[0].font.color.rgb = RGBColor(*DIM)
            p.space_after = Pt(12)
    prs.save(OUT / "nhdw-overview-deck.pptx")
    print("  nhdw-overview-deck.pptx")


def build_video() -> None:
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        sys.exit("the walkthrough needs Pillow:  pip install pillow")
    if not shutil.which("ffmpeg"):
        print("  (no ffmpeg — skipping the walkthrough)")
        return

    W, H, FPS, PER = 1280, 720, 24, 40
    BG = (14, 16, 20)

    def font(sz, bold=False):
        base = "/usr/share/fonts/truetype/dejavu/DejaVuSans"
        try:
            return ImageFont.truetype(f"{base}{'-Bold' if bold else ''}.ttf", sz)
        except OSError:
            return ImageFont.load_default()

    with tempfile.TemporaryDirectory() as tmp:
        frames = pathlib.Path(tmp)
        n = 0
        for i, (title, sub, bullets) in enumerate(SLIDES):
            for f in range(PER):
                im = Image.new("RGB", (W, H), BG)
                d = ImageDraw.Draw(im)
                d.rectangle([0, 0, W, 3], fill=(32, 36, 44))
                d.rectangle([0, 0, int(W * (i * PER + f) / (len(SLIDES) * PER)), 3], fill=(69, 207, 149))
                fade = min(1.0, (f + 3) / 8)   # frame 0 already carries the slide
                def col(c):
                    return tuple(int(BG[k] + (c[k] - BG[k]) * fade) for k in range(3))
                d.text((90, 250), title, font=font(58, True), fill=col((238, 240, 244)))
                d.text((90, 340), sub, font=font(30), fill=col((150, 158, 172)))
                d.text((90, 400), bullets[0], font=font(24), fill=col((69, 207, 149)))
                d.text((90, H - 80), f"{i + 1} / {len(SLIDES)}", font=font(20), fill=(70, 78, 92))
                im.save(frames / f"f{n:04d}.png")
                n += 1
        subprocess.run(
            ["ffmpeg", "-y", "-loglevel", "error", "-framerate", str(FPS),
             "-i", str(frames / "f%04d.png"), "-c:v", "libx264", "-pix_fmt", "yuv420p",
             "-profile:v", "baseline", "-level", "3.1", "-movflags", "+faststart",
             "-crf", "26", str(OUT / "nhdw-walkthrough.mp4")], check=True)
    print(f"  nhdw-walkthrough.mp4  ({n / FPS:.0f}s)")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    build_deck()
    build_video()
